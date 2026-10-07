#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
WG Craft - bulk entry of estimated meter readings.

Usage (processes the single .xlsx next to main.py or in excel/, or --excel):
    python main.py --dry-run           # no browser: list accounts to process
    python main.py --test              # test run: first 3 accounts
    python main.py                     # full run
    python main.py --force             # also re-check accounts done this month
    python main.py --excel list.xlsx   # use a specific file
"""

import argparse
import csv
import json
import os
import sys
import time
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from selenium import webdriver
from selenium.common.exceptions import WebDriverException
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys

BASE = Path(__file__).resolve().parent
REPORTS = BASE / "reports"
MAX_ERRORS_IN_ROW = 4

# defaults, so the script also runs without config.json
DEFAULT_CFG = {
    "start_increment": 4,
    "max_increment": None,
    "safety_stop": 500,
    "test_count": 3,
    "wait_seconds": 15,
    "step_delay": 0.8,
    "reading_date": None,
    "only_visible_rows": True,
}

# ======================================================================
#  JS helpers (find page elements by visible text and position)
# ======================================================================
JS_LIB = r"""
const vis = e => { const r = e.getBoundingClientRect(); if (r.width < 1 || r.height < 1) return false;
  const s = getComputedStyle(e); return s.visibility !== 'hidden' && s.display !== 'none'; };
const own = e => Array.from(e.childNodes).filter(n => n.nodeType === 3).map(n => n.textContent).join(' ').replace(/\s+/g, ' ').trim();
const all = () => Array.from(document.querySelectorAll('body *')).filter(vis);
const byText = (t, exact) => all().filter(e => exact ? own(e) === t : own(e).includes(t));
const onTop = e => { const r = e.getBoundingClientRect();
  const t = document.elementFromPoint((r.left + r.right) / 2, (r.top + r.bottom) / 2);
  return !!t && (t === e || e.contains(t)); };
const realInputs = () => Array.from(document.querySelectorAll('input'))
  .filter(i => vis(i) && !['hidden', 'checkbox', 'radio', 'button', 'submit'].includes(i.type));
"""

JS_TEXT = "return byText(arguments[0], arguments[1]).length > 0;"

# numbers under the "Последние показания" (last reading) column
JS_LAST = r"""
const h = byText('Последние показания', false).pop();
if (!h) return [];
const hr = h.getBoundingClientRect();
const out = all().filter(e => {
  if (!/^\d+([.,]\d+)?$/.test(own(e))) return false;
  const r = e.getBoundingClientRect(); const c = (r.left + r.right) / 2;
  return r.top >= hr.bottom - 2 && r.top - hr.bottom < 150 && c >= hr.left - 5 && c <= hr.right + 5 && onTop(e);
});
out.sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
return out.map(own);
"""

# input fields under the "Текущие показания" (current reading) column
JS_INPUTS = r"""
const h = byText('Текущие показания', false).pop();
if (!h) return [];
const hr = h.getBoundingClientRect();
const out = realInputs().filter(i => {
  const r = i.getBoundingClientRect(); const c = (r.left + r.right) / 2;
  return r.top >= hr.bottom - 2 && r.top - hr.bottom < 150 && c >= hr.left - 5 && c <= hr.right + 5 && onTop(i);
});
out.sort((a, b) => a.getBoundingClientRect().top - b.getBoundingClientRect().top);
return out;
"""

# date input next to a given label
JS_DATE = r"""
const l = byText(arguments[0], true).pop();
if (!l) return null;
const lr = l.getBoundingClientRect(); const cy = (lr.top + lr.bottom) / 2;
const c = realInputs().filter(i => { const r = i.getBoundingClientRect();
  return Math.abs((r.top + r.bottom) / 2 - cy) < 14 && r.left >= lr.right - 2 && onTop(i); });
c.sort((a, b) => a.getBoundingClientRect().left - b.getBoundingClientRect().left);
return c.length ? c[0] : null;
"""

# the "Лицевой счет" (account) field of the search form - only when the form is open
JS_SEARCH = r"""
if (!byText('Очистить', true).length) return null;
const l = byText('Лицевой счет', true).pop();
if (!l) return null;
const lr = l.getBoundingClientRect(); const cy = (lr.top + lr.bottom) / 2;
const c = realInputs().filter(i => { const r = i.getBoundingClientRect();
  return Math.abs((r.top + r.bottom) / 2 - cy) < 14 && r.left >= lr.right - 2 && onTop(i); });
c.sort((a, b) => a.getBoundingClientRect().left - b.getBoundingClientRect().left);
return c.length ? c[0] : null;
"""

# error text in the dialog (known phrases or a pink/red alert box)
JS_ERR = r"""
const phr = ['настроенного порога', 'неактивных счетчиков', 'Внос показаний невозможен', 'невозможен'];
for (const p of phr) { const m = byText(p, false); if (m.length) return own(m[0]); }
const lbl = byText('Игнорировать предыдущие показания', false)[0];
if (!lbl) return null;
let scope = document.body; let n = lbl;
while (n && n !== document.body) {
  const s = getComputedStyle(n); const r = n.getBoundingClientRect();
  if ((s.position === 'fixed' || s.position === 'absolute') && r.width > 350 && r.height > 150) { scope = n; break; }
  n = n.parentElement;
}
const rd = s => { const m = s.match(/rgba?\((\d+),\s*(\d+),\s*(\d+)(?:,\s*([\d.]+))?/);
  return m ? [+m[1], +m[2], +m[3], m[4] === undefined ? 1 : +m[4]] : null; };
const pink = e => { let n = e;
  for (let k = 0; k < 6 && n; k++) { const c = rd(getComputedStyle(n).backgroundColor);
    if (c && c[3] > 0.2) return c[0] >= 200 && c[0] - c[1] >= 25 && c[0] - c[2] >= 25; n = n.parentElement; }
  return false; };
const hit = Array.from(scope.querySelectorAll('*')).filter(e => vis(e) && own(e).length > 2 && own(e).length < 400
  && (pink(e) || /alert-danger|alert-error|text-danger|\berror\b/i.test(String(e.className))));
return hit.length ? own(hit[0]) : null;
"""

JS_ON_TOP = ("const e=arguments[0]; const r=e.getBoundingClientRect();"
             "const t=document.elementFromPoint((r.left+r.right)/2,(r.top+r.bottom)/2);"
             "return !!t && (t===e || e.contains(t));")

# ======================================================================
#  XPaths (by visible labels)
# ======================================================================
XP_ABON = ["//*[normalize-space(text())='Абоненты']"]
XP_FL = ["//*[normalize-space(text())='Абоненты ФЛ']"]
XP_TAB = ["//*[normalize-space(text())='Счетчики']"]
XP_ADD = ["//button[normalize-space(.)='Добавление показаний']",
          "//a[normalize-space(.)='Добавление показаний']",
          "//*[normalize-space(text())='Добавление показаний']"]
XP_APPLY = ["//button[normalize-space(.)='Применить']", "//*[normalize-space(text())='Применить']"]
XP_YES = ["//button[normalize-space(.)='Да']", "//a[normalize-space(.)='Да']",
          "//*[normalize-space(text())='Да']"]
XP_CANCEL = ["//button[normalize-space(.)='Отменить']", "//*[normalize-space(text())='Отменить']"]


# ======================================================================
#  Generic helpers
# ======================================================================
def js(driver, body, *args):
    return driver.execute_script(JS_LIB + body, *args)


def vis_els(driver, xp):
    out = []
    try:
        els = driver.find_elements(By.XPATH, xp)
    except WebDriverException:
        return out
    for e in els:
        try:
            if e.is_displayed():
                out.append(e)
        except WebDriverException:
            pass
    return out


def wait_for(fn, timeout, interval=0.3):
    end = time.time() + timeout
    while time.time() < end:
        try:
            r = fn()
            if r:
                return r
        except WebDriverException:
            pass
        time.sleep(interval)
    return None


def _on_top(driver, el):
    try:
        return bool(driver.execute_script(JS_ON_TOP, el))
    except WebDriverException:
        return False


def click(driver, xpaths, timeout, what="", top=False):
    """Click a visible element. top=True picks the one in the top-most dialog."""
    if isinstance(xpaths, str):
        xpaths = [xpaths]

    def find():
        for xp in xpaths:
            els = vis_els(driver, xp)
            if top and els:
                els = [e for e in els if _on_top(driver, e)] or els
            if els:
                return els[-1] if top else els[0]
        return None

    el = wait_for(find, timeout)
    if not el:
        raise TimeoutError(f"not found: {what or xpaths[0]}")
    try:
        el.click()
    except WebDriverException:
        driver.execute_script("arguments[0].click();", el)
    return el


def type_into(el, text):
    try:
        el.click()
    except WebDriverException:  # covered by another element, e.g. a placeholder
        el.parent.execute_script("arguments[0].focus();", el)
    el.send_keys(Keys.CONTROL, "a")
    el.send_keys(Keys.DELETE)
    el.send_keys(text)


def to_num(s):
    return float(str(s).replace(",", ".").replace(" ", ""))


def same_value(el, v):
    try:
        return to_num(el.get_attribute("value") or "") == float(v)
    except (ValueError, WebDriverException):
        return False


def res(status, last="", value="", inc="", note=""):
    return {"status": status, "last": last, "value": value, "inc": inc, "note": note}


# ======================================================================
#  Excel
# ======================================================================
def read_accounts(path, only_visible=True):
    """Account numbers from the 'Лиц.счет' column, skipping rows hidden by a filter."""
    ws = load_workbook(path, data_only=True).active
    col = hdr = None
    for row in ws.iter_rows(min_row=1, max_row=15):
        for c in row:
            v = str(c.value or "").lower().replace(" ", "")
            if "лиц" in v and "счет" in v and len(v) < 25:
                col, hdr = c.column, c.row
                break
        if col:
            break
    if not col:
        sys.exit("Column 'Лиц.счет' (account number) not found in the Excel file.")
    accounts, seen = [], set()
    for r in range(hdr + 1, ws.max_row + 1):
        if only_visible and ws.row_dimensions[r].hidden:
            continue
        v = ws.cell(r, col).value
        if v is None:
            continue
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        s = str(v).strip()
        if isinstance(v, int):
            s = s.zfill(10)
        if not s.isdigit() or len(s) < 8 or s in seen:
            continue
        seen.add(s)
        accounts.append(s)
    return accounts


# ======================================================================
#  Browser: start, login, navigation
# ======================================================================
def make_driver():
    opts = webdriver.ChromeOptions()
    opts.add_argument("--start-maximized")
    opts.add_experimental_option("excludeSwitches", ["enable-automation"])
    opts.add_experimental_option("prefs", {
        "credentials_enable_service": False,
        "profile.password_manager_enabled": False,
        "autofill.profile_enabled": False,
    })
    return webdriver.Chrome(options=opts)


def do_login(driver, cfg, login, pwd):
    W = cfg["wait_seconds"]
    driver.get(cfg["login_url"])

    def fields():
        ins = [e for e in vis_els(driver, "//input")
               if (e.get_attribute("type") or "text") in ("text", "password", "email")]
        return ins if len(ins) >= 2 else None

    ins = wait_for(fields, W)
    if not ins:
        raise TimeoutError("login fields not found")
    type_into(ins[0], login)
    type_into(ins[1], pwd)
    btns = vis_els(driver, "//button[normalize-space(.)='Вход']") or vis_els(driver, "//input[@type='submit']")
    if btns:
        try:
            btns[0].click()
        except WebDriverException:
            driver.execute_script("arguments[0].click();", btns[0])
    else:
        ins[1].send_keys(Keys.ENTER)
    if not wait_for(lambda: "/login" not in driver.current_url, W * 2):
        raise RuntimeError("Login failed - check the credentials in .env")


def search_input(driver):
    """Only the account field of the search form (never any other input)."""
    return js(driver, JS_SEARCH)


def open_subscribers_menu(driver, cfg):
    W = cfg["wait_seconds"]
    if not vis_els(driver, XP_FL[0]):
        for el in vis_els(driver, XP_ABON[0]):
            try:
                el.click()
            except WebDriverException:
                driver.execute_script("arguments[0].click();", el)
            if wait_for(lambda: vis_els(driver, XP_FL[0]), 3):
                break
    click(driver, XP_FL, W, what="Абоненты ФЛ")


def open_search(driver, cfg):
    """Make sure the search form is open (via the menu or by reloading the page)."""
    W = cfg["wait_seconds"]
    el = search_input(driver)
    if el:
        return el
    try:
        open_subscribers_menu(driver, cfg)
        el = wait_for(lambda: search_input(driver), 5)
        if el:
            return el
    except TimeoutError:
        pass
    driver.get(cfg["subscribers_url"])
    el = wait_for(lambda: search_input(driver), W)
    if not el:
        open_subscribers_menu(driver, cfg)
        el = wait_for(lambda: search_input(driver), W)
    if not el:
        raise TimeoutError("search field not found")
    return el


# ======================================================================
#  Dialog handling
# ======================================================================
def modal_open(driver):
    return js(driver, JS_TEXT, "Игнорировать предыдущие показания", False)


def get_error(driver):
    return js(driver, JS_ERR)


def cancel_modal(driver, timeout=6):
    try:
        click(driver, XP_CANCEL, timeout, what="Отменить", top=True)
        wait_for(lambda: not modal_open(driver), timeout)
    except (TimeoutError, WebDriverException):
        pass


def classify(err):
    low = err.lower()
    if "порога" in low:  # daily-average threshold (either wording): retry with +1
        return "retry"
    if "счетчик" in low:  # meter removed / inactive: skip the account
        return "meter"
    return "other"


def ensure_date(driver, want):
    el = js(driver, JS_DATE, "Дата показаний")
    if el is None:
        return
    if (el.get_attribute("value") or "").strip() == want:
        return
    type_into(el, want)
    el.send_keys(Keys.TAB)


def process(driver, cfg, acc, rdate):
    W = cfg["wait_seconds"]
    pause = cfg["step_delay"]

    inp = open_search(driver, cfg)
    type_into(inp, acc)
    if (inp.get_attribute("value") or "").strip() != acc:
        raise RuntimeError("account number was not typed into the search field correctly")
    inp.send_keys(Keys.ENTER)
    if not wait_for(lambda: js(driver, JS_TEXT, acc, True), W * 1.5):
        return res("not_found", note="subscriber card did not open")
    time.sleep(pause)
    click(driver, XP_TAB, W, what="Счетчики tab")
    time.sleep(pause)

    inc = cfg["start_increment"]
    max_inc = cfg.get("max_increment")
    safety = cfg.get("safety_stop", 500)

    while True:
        try:
            click(driver, XP_ADD, W * 2, what="Добавление показаний")
        except TimeoutError:
            return res("skip_other", note="'Add reading' button not found "
                                        "(no meter?) - check manually")
        if not wait_for(lambda: modal_open(driver) or get_error(driver), W):
            raise TimeoutError("'Add reading' dialog did not open")
        time.sleep(pause)

        # an error shown as soon as the dialog opens (e.g. the meter was removed)
        err = get_error(driver)
        if err:
            cancel_modal(driver)
            kind = classify(err)
            return res("skip_meter" if kind == "meter" else "skip_other", note=err)

        last_vals = js(driver, JS_LAST)
        inputs = js(driver, JS_INPUTS)
        if len(last_vals) != 1 or len(inputs) != 1:
            cancel_modal(driver)
            return res("skip_other",
                       note=f"could not identify a single meter (last={len(last_vals)}, inputs={len(inputs)})")
        base = to_num(last_vals[0])
        if base != int(base):
            cancel_modal(driver)
            return res("skip_other", last=last_vals[0], note="fractional reading - enter manually")
        base = int(base)

        # already entered today - do not enter it twice
        ld = js(driver, JS_DATE, "Дата посл. показания")
        if ld is not None and (ld.get_attribute("value") or "").strip() == rdate:
            cancel_modal(driver)
            return res("already", last=base, note="entered today")

        value = base + inc
        ensure_date(driver, rdate)
        field = inputs[0]
        type_into(field, str(value))
        field.send_keys(Keys.TAB)
        try:
            if not same_value(field, value):
                driver.execute_script(
                    "arguments[0].value=arguments[1];"
                    "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));"
                    "arguments[0].dispatchEvent(new Event('change',{bubbles:true}));",
                    field, str(value))
        except WebDriverException:
            pass

        click(driver, XP_APPLY, W, what="Применить", top=True)
        stage = wait_for(lambda: "confirm" if any(vis_els(driver, x) for x in XP_YES)
                         else ("err" if get_error(driver) else None), W)
        if stage == "confirm":
            click(driver, XP_YES, W, what="Да", top=True)

        def state():
            e = get_error(driver)
            if e:
                return ("err", e)
            if not modal_open(driver):
                return ("closed", "")
            lv = js(driver, JS_LAST)
            if len(lv) == 1 and to_num(lv[0]) == value:
                return ("updated", "")
            return None

        st = wait_for(state, W * 1.5)
        if st is None:
            cancel_modal(driver)
            return res("error", last=base, value=value, inc=inc,
                       note="no response - check on the site manually")
        kind, txt = st
        if kind in ("closed", "updated"):
            if kind == "updated":
                cancel_modal(driver)
            return res("ok", last=base, value=value, inc=inc)

        k = classify(txt)
        cancel_modal(driver)
        if k == "retry":
            inc += 1
            if (max_inc and inc > max_inc) or inc > safety:
                return res("skip_threshold", last=base, value=value, inc=inc - 1, note="increment limit reached")
            continue
        if k == "meter":
            return res("skip_meter", last=base, note=txt)
        return res("skip_other", last=base, value=value, inc=inc, note=txt)


# ======================================================================
#  Reporting and resuming
# ======================================================================
def recover(driver, cfg, login, pwd):
    """Bring the browser back to a working state after an error (lost connection, expired session)."""
    for _ in range(3):
        try:
            driver.get(cfg["subscribers_url"])
            time.sleep(3)
            if "/login" in driver.current_url:
                do_login(driver, cfg, login, pwd)
                driver.get(cfg["subscribers_url"])
            ok = wait_for(lambda: vis_els(driver, XP_ABON[0]) or vis_els(driver, XP_FL[0])
                          or search_input(driver), cfg["wait_seconds"])
            if ok:
                return True
        except Exception:  # noqa: BLE001
            pass
        time.sleep(10)
    return False


def dump_debug(driver, acc):
    d = REPORTS / "debug"
    d.mkdir(parents=True, exist_ok=True)
    try:
        driver.save_screenshot(str(d / f"{acc}.png"))
        (d / f"{acc}.html").write_text(driver.page_source, encoding="utf-8")
    except Exception:
        pass


def load_done(path):
    return set(path.read_text(encoding="utf-8").split()) if path.exists() else set()


def find_excel():
    """The single .xlsx next to main.py or in excel/ (any file name)."""
    found = [p for d in (BASE / "excel", BASE) for p in sorted(d.glob("*.xlsx"))
             if not p.name.startswith("~$")]
    if len(found) == 1:
        return found[0]
    if not found:
        sys.exit("No Excel file found. Put the filtered .xlsx next to main.py or use --excel.")
    sys.exit("Several .xlsx files found: " + ", ".join(p.name for p in found)
             + ". Keep only one, or pass --excel.")


# ======================================================================
def main():
    ap = argparse.ArgumentParser(description="Enter estimated meter readings for a list of accounts.")
    ap.add_argument("--excel", help="path to the Excel file (default: the single .xlsx found)")
    ap.add_argument("--test", action="store_true", help="only the first test_count accounts")
    ap.add_argument("--force", action="store_true", help="also re-check accounts already done this month")
    ap.add_argument("--dry-run", action="store_true", help="no browser: only list the accounts to process")
    args = ap.parse_args()

    cfg = dict(DEFAULT_CFG)
    cfg_path = BASE / "config.json"
    if cfg_path.exists():  # overrides the defaults
        with open(cfg_path, encoding="utf-8") as f:
            cfg.update(json.load(f))

    xl = Path(args.excel) if args.excel else find_excel()
    if not xl.exists():
        sys.exit(f"Excel file not found: {xl}")
    accounts = read_accounts(xl, cfg.get("only_visible_rows", True))

    now = datetime.now()
    rdate = cfg.get("reading_date") or now.strftime("%d.%m.%Y")
    REPORTS.mkdir(exist_ok=True)
    done_path = REPORTS / f"done_{now:%Y-%m}.txt"
    done = set() if args.force else load_done(done_path)
    todo = [a for a in accounts if a not in done]
    if args.test:
        todo = todo[: cfg.get("test_count", 3)]

    print(f"File: {xl.name} | In Excel: {len(accounts)} | Done this month: {len(done)} | "
          f"To process: {len(todo)} | Start increment: +{cfg['start_increment']}")
    if not todo:
        print("No accounts left to process.")
        return
    if args.dry_run:
        for acc in todo[:50]:
            print("  ", acc)
        if len(todo) > 50:
            print(f"   ... and {len(todo) - 50} more")
        return

    load_dotenv(BASE / ".env")
    base_url = (os.getenv("WG_BASE_URL") or "").rstrip("/")
    login, pwd = os.getenv("WG_LOGIN"), os.getenv("WG_PASSWORD")
    if not base_url or not login or not pwd:
        sys.exit(".env must contain WG_BASE_URL, WG_LOGIN and WG_PASSWORD (see .env.example).")
    cfg.setdefault("login_url", base_url + "/login")
    cfg.setdefault("subscribers_url", base_url + "/fcmf")

    report_path = REPORTS / f"report_{now:%Y%m%d_%H%M}.csv"
    driver = make_driver()
    errors_in_row = 0
    try:
        do_login(driver, cfg, login, pwd)
        with open(report_path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(["account", "status", "last_reading", "entered", "increment", "message"])
            for i, acc in enumerate(todo, 1):
                try:
                    if "/login" in driver.current_url:
                        do_login(driver, cfg, login, pwd)
                    r = process(driver, cfg, acc, rdate)
                    errors_in_row = 0
                except Exception as e:  # noqa: BLE001
                    errors_in_row += 1
                    dump_debug(driver, acc)
                    r = res("error", note=f"{type(e).__name__}: {str(e)[:150]}")
                    try:
                        cancel_modal(driver, 3)
                    except Exception:  # noqa: BLE001
                        pass
                    print("  recovering the browser...")
                    recover(driver, cfg, login, pwd)
                w.writerow([acc, r["status"], r["last"], r["value"], r["inc"], r["note"]])
                fh.flush()
                if r["status"] in ("ok", "already"):
                    with open(done_path, "a", encoding="utf-8") as df:
                        df.write(acc + "\n")
                step = f"+{r['inc']}" if r["inc"] != "" else ""
                print(f"[{i}/{len(todo)}] {acc}  {r['status']:<15} {r['last']}->{r['value']} {step} {r['note']}")
                if errors_in_row >= MAX_ERRORS_IN_ROW:
                    print(f"\n{MAX_ERRORS_IN_ROW} errors in a row - stopped. "
                          "See the screenshots and HTML in reports/debug.")
                    break
    except KeyboardInterrupt:
        print("\nStopped (Ctrl+C). Run again to resume.")
    finally:
        try:
            driver.quit()
        except Exception:  # noqa: BLE001
            pass
        print(f"\nReport: {report_path}")


if __name__ == "__main__":
    main()
